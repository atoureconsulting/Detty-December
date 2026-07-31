"""Render a pptx to HTML pages by reading real shape geometry, for visual QA."""
import sys, html, os
from pptx import Presentation
from pptx.util import Emu

SCALE = 96  # px per inch

def emu_in(v):
    return v / 914400.0 if v is not None else 0.0

def solid(fmt):
    try:
        if fmt.type is not None and fmt.type == 1:  # MSO_FILL.SOLID
            c = fmt.fore_color
            if c.type == 1:  # RGB
                return "#" + str(c.rgb)
    except Exception:
        pass
    return None

def line_color(shp):
    try:
        ln = shp.line
        if ln.color and ln.color.type == 1:
            return "#" + str(ln.color.rgb), max(emu_in(ln.width) * SCALE, 0.7) if ln.width else 1
    except Exception:
        pass
    return None, 0

def render(path, out_dir):
    prs = Presentation(path)
    SW = emu_in(prs.slide_width) * SCALE
    SH = emu_in(prs.slide_height) * SCALE
    os.makedirs(out_dir, exist_ok=True)
    pages = []
    for idx, slide in enumerate(prs.slides, 1):
        bg = "#FFFFFF"
        try:
            b = solid(slide.background.fill)
            if b: bg = b
        except Exception:
            pass
        parts = []
        for shp in slide.shapes:
            x, y = emu_in(shp.left) * SCALE, emu_in(shp.top) * SCALE
            w, h = emu_in(shp.width) * SCALE, emu_in(shp.height) * SCALE
            fill = None
            try:
                fill = solid(shp.fill)
            except Exception:
                pass
            lc, lw = line_color(shp)
            prst = ""
            try:
                A = "{http://schemas.openxmlformats.org/drawingml/2006/main}"
                g = shp._element.find(f".//{A}prstGeom")
                if g is not None:
                    prst = g.get("prst") or ""
            except Exception:
                pass
            radius = ""
            if prst == "ellipse":
                radius = "border-radius:50%;"
            elif prst == "roundRect":
                radius = "border-radius:6px;"
            style = (f"position:absolute;left:{x:.2f}px;top:{y:.2f}px;"
                     f"width:{w:.2f}px;height:{h:.2f}px;{radius}")
            if fill: style += f"background:{fill};"
            if lc: style += f"border:{lw:.2f}px solid {lc};box-sizing:border-box;"
            if not shp.has_text_frame:
                parts.append(f'<div style="{style}"></div>')
                continue
            tf = shp.text_frame
            # vertical anchor
            anchor = "flex-start"
            try:
                a = str(tf.vertical_anchor)
                if "MIDDLE" in a: anchor = "center"
                elif "BOTTOM" in a: anchor = "flex-end"
            except Exception:
                pass
            style += f"display:flex;flex-direction:column;justify-content:{anchor};overflow:visible;"
            inner = []
            for para in tf.paragraphs:
                al = "left"
                try:
                    pa = str(para.alignment)
                    if "CENTER" in pa: al = "center"
                    elif "RIGHT" in pa: al = "right"
                except Exception:
                    pass
                runs = []
                fs, fc, fb, fi, ff, cs = 18, "#000000", False, False, "Calibri", 0
                for r in para.runs:
                    f = r.font
                    if f.size: fs = f.size.pt
                    if f.bold: fb = True
                    if f.italic: fi = True
                    if f.name: ff = f.name
                    try:
                        if f.color and f.color.type == 1: fc = "#" + str(f.color.rgb)
                    except Exception:
                        pass
                    runs.append(html.escape(r.text))
                txt = "".join(runs)
                if not txt.strip():
                    continue
                bullet = ""
                ind = ""
                # crude bullet detection via paragraph XML
                if para._pPr is not None and para._pPr.find(
                        '{http://schemas.openxmlformats.org/drawingml/2006/main}buChar') is not None:
                    bullet = "• "
                    ind = "padding-left:11px;text-indent:-11px;"
                lh = fs * 1.22
                inner.append(
                    f'<div style="font-family:\'{ff}\',sans-serif;font-size:{fs}px;'
                    f'color:{fc};font-weight:{"700" if fb else "400"};'
                    f'font-style:{"italic" if fi else "normal"};text-align:{al};'
                    f'line-height:{lh:.1f}px;margin:0 0 2px 0;{ind}">{bullet}{txt}</div>')
            parts.append(f'<div style="{style}">{"".join(inner)}</div>')
        pages.append(
            f'<div class="slide" style="width:{SW:.0f}px;height:{SH:.0f}px;background:{bg};">'
            f'{"".join(parts)}'
            f'<div class="num">{idx}</div></div>')
    doc = ("<style>body{margin:0;background:#555;font-family:Calibri,sans-serif;}"
           ".slide{position:relative;margin:0 auto 16px;overflow:hidden;}"
           ".num{position:absolute;right:6px;bottom:4px;font-size:11px;color:#888;}"
           "</style>" + "".join(pages))
    with open(os.path.join(out_dir, "preview.html"), "w") as f:
        f.write(doc)
    print(f"{len(pages)} slides -> {out_dir}/preview.html  ({SW:.0f}x{SH:.0f})")

if __name__ == "__main__":
    render(sys.argv[1], sys.argv[2])
